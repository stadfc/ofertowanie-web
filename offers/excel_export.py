from io import BytesIO
from urllib.parse import quote
from urllib.request import Request, urlopen

from django.conf import settings
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

PHOTO_W = 110
PHOTO_H = 72


def _thin():
    side = Side(style="thin", color="B4B4B4")
    return Border(left=side, right=side, top=side, bottom=side)


def _fetch_photo(filename: str) -> bytes | None:
    name = (filename or "").replace("\\", "/").split("/")[-1]
    if not name or ".." in name:
        return None
    url = settings.PHOTO_BASE_URL.rstrip("/") + "/" + quote(name)
    req = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urlopen(req, timeout=15) as resp:
            return resp.read()
    except Exception:
        return None


def _xl_image(raw: bytes) -> XLImage | None:
    try:
        src = PILImage.open(BytesIO(raw))
        src = src.convert("RGBA")
        src.thumbnail((PHOTO_W, PHOTO_H), PILImage.Resampling.LANCZOS)
        canvas = PILImage.new("RGBA", (PHOTO_W, PHOTO_H), (255, 255, 255, 0))
        x = (PHOTO_W - src.width) // 2
        y = (PHOTO_H - src.height) // 2
        canvas.paste(src, (x, y), src)
        buf = BytesIO()
        canvas.convert("RGB").save(buf, format="PNG")
        buf.seek(0)
        buf.name = "photo.png"
        img = XLImage(buf)
        img.width = PHOTO_W
        img.height = PHOTO_H
        return img
    except Exception:
        return None


def build_offer_xlsx(offer, lines, profile, handlowiec, valid_to, metrics) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Oferta"
    accent = "2F6B4F"
    head_fill = PatternFill("solid", fgColor="D0D0D0")
    title_font = Font(name="Calibri", size=16, bold=True, color=accent)
    label_font = Font(name="Calibri", size=10, bold=True, color="666666")
    body_font = Font(name="Calibri", size=11)
    header_font = Font(name="Calibri", size=11, bold=True)
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    money = Alignment(horizontal="right", vertical="center")
    border = _thin()
    with_photos = bool(offer.with_photos)

    ws.merge_cells("A1:G1")
    ws["A1"] = offer.nazwa or "Propozycja"
    ws["A1"].font = title_font
    ws.row_dimensions[1].height = 22

    ws.merge_cells("A2:G2")
    meta = f"{handlowiec}"
    if profile.stanowisko:
        meta += f" · {profile.stanowisko}"
    meta += f" · {offer.created_on.strftime('%d-%m-%Y')} — {valid_to.strftime('%d-%m-%Y')}"
    ws["A2"] = meta
    ws["A2"].font = body_font

    contact_bits = []
    if profile.telefon:
        contact_bits.append(f"tel. {profile.telefon}")
    if profile.email:
        contact_bits.append(profile.email)
    if profile.kontakt:
        contact_bits.append(profile.kontakt)
    ws.merge_cells("A3:G3")
    ws["A3"] = " · ".join(contact_bits)
    ws["A3"].font = Font(name="Calibri", size=10, color="555555")

    row = 5
    if profile.komunikat_otwarcia:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row, 1, profile.komunikat_otwarcia).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row, 1).font = body_font
        ws.row_dimensions[row].height = 48
        row += 2

    ws.cell(row, 1, "Płatnik").font = label_font
    ws.cell(row, 4, "Odbiorca").font = label_font
    row += 1
    platnik = "\n".join(
        x for x in [offer.platnik_nazwa, offer.platnik_adres, f"NIP {offer.platnik_nip}" if offer.platnik_nip else ""] if x
    )
    odbiorca = "\n".join(x for x in [offer.odbiorca_nazwa, offer.odbiorca_adres] if x)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=7)
    ws.cell(row, 1, platnik).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row, 4, odbiorca).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 42
    row += 2

    headers = []
    if with_photos:
        headers.append("Zdjęcie")
    headers.extend(
        ["Opis", "Cena katalogowa netto", "Rabat %", "Cena netto dla Ciebie", "Ilość", "Wartość netto"]
    )
    header_row = row
    for col, title in enumerate(headers, 1):
        cell = ws.cell(header_row, col, title)
        cell.fill = head_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    row += 1

    def num(value):
        if value in (None, ""):
            return None
        try:
            return float(str(value).replace(",", "."))
        except ValueError:
            return None

    first_data = row
    holds = []
    for line in lines:
        if with_photos:
            ws.row_dimensions[row].height = 58
            cell = ws.cell(row, 1, "")
            cell.border = border
            cell.alignment = center
            raw = _fetch_photo(line.get("zdjecie_plik") or "")
            if raw:
                img = _xl_image(raw)
                if img:
                    holds.append(img)
                    img.anchor = f"A{row}"
                    ws.add_image(img)
        values = [
            line.get("opis") or "",
            num(line.get("cena")),
            num(line.get("rabat")),
            num(line.get("netto")),
            num(line.get("ilosc")),
            num(line.get("wartosc")),
        ]
        start = 2 if with_photos else 1
        for offset, value in enumerate(values):
            cell = ws.cell(row, start + offset, value)
            cell.font = body_font
            cell.border = border
            cell.alignment = wrap if offset == 0 else money
            if offset == 0:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            if offset >= 1 and isinstance(value, float):
                cell.number_format = '#,##0.00'
        row += 1

    last_data = row - 1
    if last_data >= first_data:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_data}"

    total_font = Font(name="Calibri", size=11, bold=True)
    sum_fill = PatternFill("solid", fgColor="F3F3F3")
    last_col = len(headers)
    opis_col = 2 if with_photos else 1

    def summary_row(label, amount):
        nonlocal row
        for col in range(1, last_col + 1):
            cell = ws.cell(row, col, None)
            cell.border = border
            cell.fill = sum_fill
            cell.font = total_font
        ws.cell(row, opis_col, label).alignment = Alignment(vertical="center")
        amount_cell = ws.cell(row, last_col, num(amount))
        if isinstance(amount_cell.value, float):
            amount_cell.number_format = '#,##0.00'
            amount_cell.alignment = money
        row += 1

    summary_row("Transport netto", metrics.get("transport") or "0")
    summary_row("Razem netto", metrics.get("wartosc") or "0")
    row += 1
    if profile.komunikat_zamkniecia:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        ws.cell(row, 1, profile.komunikat_zamkniecia).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 42
        row += 2
    footer = handlowiec
    if profile.stanowisko:
        footer += f", {profile.stanowisko}"
    extras = [x for x in [profile.telefon, profile.email] if x]
    if extras:
        footer += " · " + " · ".join(extras)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    ws.cell(row, 1, footer).font = Font(name="Calibri", size=10)

    widths = [16, 48, 22, 12, 22, 10, 16] if with_photos else [48, 22, 12, 22, 10, 16]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{header_row + 1}"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
