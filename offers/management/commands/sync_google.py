import io
from decimal import Decimal, InvalidOperation
from urllib.request import Request, urlopen

from django.conf import settings
from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from offers.models import Product


class Command(BaseCommand):
    help = "Pobiera część katalogu BAZA z opublikowanego arkusza Google (wiersze ze zdjęciem)."

    def add_arguments(self, parser):
        parser.add_argument("--limit", type=int, default=250)
        parser.add_argument("--url", default=settings.GOOGLE_BAZA_XLSX)

    def handle(self, *args, **options):
        url = options["url"]
        limit = options["limit"]
        self.stdout.write(f"Pobieranie {url}")
        req = Request(url, headers={"User-Agent": "OfertowanieWeb/0.1"})
        with urlopen(req, timeout=90) as resp:
            data = resp.read()
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(rows)]
        idx = {name: i for i, name in enumerate(header)}
        for col in ("Kod", "Nazwa", "Combi", "Cena_det", "Zdjecie"):
            if col not in idx:
                self.stderr.write(f"Brak kolumny {col}: {header}")
                return

        Product.objects.all().delete()
        created = 0
        scanned = 0
        for row in rows:
            scanned += 1
            zdj = str(row[idx["Zdjecie"]] or "").strip()
            if not zdj:
                continue
            kod = str(row[idx["Kod"]] or "").strip()
            nazwa = str(row[idx["Nazwa"]] or "").strip()
            combi = str(row[idx["Combi"]] or "").strip() or f"{kod} {nazwa}".strip()
            raw_cena = row[idx["Cena_det"]]
            try:
                cena = Decimal(str(raw_cena).replace(",", ".")) if raw_cena not in (None, "") else Decimal("0")
            except (InvalidOperation, ValueError):
                cena = Decimal("0")
            if not kod:
                continue
            raw_baz = row[idx["Cena_baz"]] if "Cena_baz" in idx else 0
            try:
                baz = Decimal(str(raw_baz).replace(",", ".")) if raw_baz not in (None, "") else Decimal("0")
            except (InvalidOperation, ValueError):
                baz = Decimal("0")
            waluta = str(row[idx["Waluta"]] or "PLN").strip() if "Waluta" in idx else "PLN"
            Product.objects.create(
                kod=kod[:64],
                nazwa=nazwa[:800],
                combi=combi[:900],
                cena_det=cena,
                cena_baz=baz,
                waluta=(waluta or "PLN")[:8],
                zdjecie=zdj[:255],
            )
            created += 1
            if created >= limit:
                break
        Product.objects.update_or_create(
            kod="TRANSPORT",
            defaults={
                "nazwa": "Koszt transportu",
                "combi": "TRANSPORT Koszt transportu",
                "cena_det": Decimal("20.00"),
                "cena_baz": Decimal("20.00"),
                "waluta": "PLN",
                "zdjecie": "",
            },
        )
        wb.close()
        self.stdout.write(self.style.SUCCESS(f"Zapisano {created} produktów ze zdjęciem (przeszukano {scanned} wierszy)."))
        first = Product.objects.exclude(zdjecie="").first()
        from offers.models import Offer, OfferLine
        offer = Offer.objects.order_by("id").first()
        if first and offer:
            line = offer.lines.order_by("position").first()
            if line:
                line.combi = first.combi
                line.rabat = Decimal("0.10")
                line.ilosc = Decimal("1")
                line.save()
                self.stdout.write(f"Przykład na ofercie: {first.combi}")
